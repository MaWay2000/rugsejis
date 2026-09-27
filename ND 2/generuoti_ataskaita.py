from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "ANFIS_rezultatai"
FIGURES = RESULTS / "grafikai"
OUTPUT = ROOT / "ataskaita"
DOCX_PATH = OUTPUT / "ND2_ANFIS_Mantas_Matusevicius_DISfm-26.docx"

NAVY = "244F75"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
TEXT = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start),
                          ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, numeric_columns: set[int] | None = None) -> None:
    numeric_columns = numeric_columns or set()
    table.autofit = True
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if col_index in numeric_columns:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
                    run.font.size = Pt(9.5)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.25)
    paragraph.add_run(text)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(9)


def add_figure(doc: Document, path: Path, caption: str, width: float) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def read_results() -> list[dict[str, str]]:
    with (RESULTS / "rezultatu_palyginimas.csv").open(
        "r", encoding="utf-8-sig", newline=""
    ) as stream:
        return list(csv.DictReader(stream))


def read_forecast(name: str) -> list[dict[str, str]]:
    with (RESULTS / name).open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def decimal(value: str | float, digits: int = 4) -> str:
    return f"{float(value):.{digits}f}".replace(".", ",")


def build_report() -> Path:
    rows = read_results()
    without_previous = rows[0]
    with_previous = rows[1]
    year_forecast = read_forecast("metu_prognoze.csv")
    forecast_2030 = read_forecast("prognoze_iki_2030.csv")
    forecast_column = "Prognozuota_kaina_EUR_uz_litra"
    year_values = [float(row[forecast_column]) for row in year_forecast]
    long_values = [float(row[forecast_column]) for row in forecast_2030]

    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = TEXT
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 12.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = TEXT
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)
    title.add_run("Benzino kainos prognozavimas naudojant ANFIS")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(18)
    subtitle_run = subtitle.add_run("Mokslinio straipsnio eksperimentų atkūrimas")
    subtitle_run.font.size = Pt(15)
    subtitle_run.bold = True

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_before = Pt(72)
    author.add_run("Studentas Mantas Matusevičius\n").bold = True
    author.add_run("Grupė DISfm-26")

    work = doc.add_paragraph()
    work.alignment = WD_ALIGN_PARAGRAPH.CENTER
    work.paragraph_format.space_before = Pt(70)
    work.add_run("Namų darbas Nr 2").bold = True
    doc.add_page_break()

    add_heading(doc, "Darbo santrauka")
    add_body(
        doc,
        "Darbe atkartoti benzino kainų prognozavimo eksperimentai, kuriuose taikomas "
        "adaptyvus neuro fuzzy išvadų modelis ANFIS. Naudoti oficialūs savaitiniai JAV "
        "benzino kainų duomenys. Geriausiai veikė modelis su ankstesnės savaitės kaina: "
        f"atkurtas RMSE buvo {decimal(with_previous['Atkartotas_RMSE'], 4)}, R2 "
        f"{decimal(with_previous['Atkartotas_R2'], 4)}, o koreliacija "
        f"{decimal(with_previous['Atkartota_koreliacija'], 4)}. Šios reikšmės yra labai "
        "artimos straipsnyje paskelbtiems RMSE, R2 ir koreliacijos rezultatams. Papildomai "
        "sudarytos 52 savaičių ir savaitinė prognozė iki 2030 m., pateikta EUR už litrą."
    )

    add_heading(doc, "1 Pasirinktas straipsnis")
    add_body(
        doc,
        "Pasirinktas E H I Eliwa, A M El Koshiry, T Abd El Hafeez ir A Omar straipsnis "
        "Optimal Gasoline Price Predictions Leveraging the ANFIS Regression Model. "
        "Jis paskelbtas žurnale International Journal of Intelligent Systems, straipsnio "
        "numeris 8462056, DOI 10.1155/2024/8462056. Autoriai nagrinėjo savaitinių benzino "
        "kainų prognozavimą ir lygino ANFIS variantus be ankstesnės kainos bei su ankstesnės "
        "savaitės kaina."
    )
    add_body(
        doc,
        "Ši tema skiriasi nuo ND1 skaitmenų klasifikavimo užduoties. Čia sprendžiama "
        "regresijos ir laiko eilutės prognozavimo užduotis, o modelis jungia neuroninį "
        "mokymąsi su fuzzy logikos narystės funkcijomis ir Sugeno taisyklėmis."
    )

    add_heading(doc, "2 Duomenys ir paruošimas")
    add_body(
        doc,
        "Duomenys gauti iš JAV Energetikos informacijos administracijos EIA. Naudota "
        "savaitinė visų markių ir visų mišinių mažmeninė benzino kaina JAV doleriais už "
        "galoną. Straipsnio eksperimentui palikta jo laikotarpio dalis su 1 583 įrašais, o "
        "ateities prognozei panaudota visa tuo metu prieinama EIA istorija. Grafikuose kainos "
        "perskaičiuotos į EUR už litrą pagal fiksuotą kursą 1 EUR = 1,1403 USD ir santykį "
        "1 JAV galonas = 3,785411784 litro."
    )
    add_figure(
        doc,
        FIGURES / "benzino_kainu_laiko_eilute.png",
        "1 pav Savaitinė JAV benzino kaina",
        6.65,
    )
    add_body(
        doc,
        "Straipsnis nurodo 1 583 įrašus, tačiau paskelbtos mokymo ir testavimo imtys sudaro "
        "tik 1 502 įrašus. Kad imčių dydžiai sutaptų, atkūrime naudoti pirmieji 1 502 įrašai. "
        "Fiksuota atsitiktinė seka padalijo juos į 1 051 mokymo ir 451 testavimo pavyzdį. "
        "Požymiai buvo diena, mėnuo ir metai. Antrajame bandyme pridėta ankstesnės savaitės "
        "kaina. Požymiai normalizuoti pagal mokymo imties minimumą ir maksimumą."
    )

    add_heading(doc, "3 ANFIS metodika")
    add_body(
        doc,
        "Sukurtas pirmos eilės Sugeno ANFIS modelis. Kiekvienam įvesties požymiui naudotos "
        "dvi Gauso narystės funkcijos. Modelis be ankstesnės kainos turėjo 8 taisykles, o "
        "modelis su ankstesne kaina turėjo 16 taisyklių. Kiekvienos taisyklės išvestis buvo "
        "tiesinė įvesties požymių funkcija."
    )
    add_bullet(doc, "Narystės funkcijų centrai ir pločiai mokyti Adam optimizatoriumi")
    add_bullet(doc, "Sugeno taisyklių išvadų koeficientai rasti mažiausių kvadratų metodu")
    add_bullet(doc, "Abu modeliai mokyti 100 epochų su ta pačia mokymo ir testavimo imtimi")
    add_bullet(doc, "Vertinimui naudoti MSE, RMSE, MAE, R2 ir koreliacijos koeficientas")
    add_body(
        doc,
        "Vietinėje MATLAB R2026a aplinkoje nebuvo įdiegtas Fuzzy Logic Toolbox, todėl "
        "ANFIS mokymas realizuotas savarankiškai baziniu MATLAB kodu. Tai leido paleisti "
        "visą eksperimentą be papildomų bibliotekų. Straipsnio Python realizacijos kodas "
        "nebuvo pateiktas, todėl tiksliai atkartoti visas vidines optimizavimo nuostatas "
        "nebuvo įmanoma."
    )
    add_figure(
        doc,
        FIGURES / "narystes_funkcijos.png",
        "2 pav Išmoktos modelio su ankstesne kaina narystės funkcijos",
        6.35,
    )

    doc.add_page_break()
    add_heading(doc, "4 Programinė ir aparatinė aplinka")
    environment = doc.add_table(rows=1, cols=2)
    environment.rows[0].cells[0].text = "Komponentas"
    environment.rows[0].cells[1].text = "Naudota aplinka"
    environment_rows = [
        ("Operacinė sistema", "Microsoft Windows 10 Pro"),
        ("Procesorius", "Intel Core i7 7700K 4,20 GHz"),
        ("Atmintis", "32 GB RAM"),
        ("MATLAB", "R2026a Update 5"),
        ("Papildomi priedai", "Nenaudoti"),
        ("Atsitiktinė seka", "42"),
        ("Mokymo epochos", "100"),
    ]
    for left, right in environment_rows:
        cells = environment.add_row().cells
        cells[0].text = left
        cells[1].text = right
    style_table(environment)
    add_caption(doc, "1 lentelė Eksperimento aplinka")
    add_body(
        doc,
        "AI agentas naudotas straipsnio paieškai ir analizei, oficialių duomenų šaltinio "
        "parinkimui, MATLAB kodo parengimui, vietiniam vykdymui, rezultatų patikrai ir "
        "ataskaitos sudarymui. Metrikos apskaičiuotos programoje iš išsaugotų testavimo "
        "prognozių."
    )

    add_heading(doc, "5 Eksperimentų rezultatai")
    comparison = doc.add_table(rows=1, cols=5)
    headers = ["Modelis", "Rodiklis", "Straipsnis", "Atkartota", "Skirtumas"]
    for cell, value in zip(comparison.rows[0].cells, headers):
        cell.text = value
    metrics = [
        ("Be ankstesnės kainos", "RMSE", without_previous["Straipsnio_RMSE"], without_previous["Atkartotas_RMSE"]),
        ("Be ankstesnės kainos", "R2", without_previous["Straipsnio_R2"], without_previous["Atkartotas_R2"]),
        ("Be ankstesnės kainos", "Koreliacija", without_previous["Straipsnio_koreliacija"], without_previous["Atkartota_koreliacija"]),
        ("Su ankstesne kaina", "RMSE", with_previous["Straipsnio_RMSE"], with_previous["Atkartotas_RMSE"]),
        ("Su ankstesne kaina", "R2", with_previous["Straipsnio_R2"], with_previous["Atkartotas_R2"]),
        ("Su ankstesne kaina", "Koreliacija", with_previous["Straipsnio_koreliacija"], with_previous["Atkartota_koreliacija"]),
    ]
    for model_name, metric, article, reproduced in metrics:
        cells = comparison.add_row().cells
        values = [
            model_name,
            metric,
            decimal(article, 4),
            decimal(reproduced, 4),
            decimal(float(reproduced) - float(article), 4),
        ]
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(comparison, numeric_columns={2, 3, 4})
    add_caption(doc, "2 lentelė Rezultatų palyginimas su straipsniu")
    add_body(
        doc,
        f"Be ankstesnės kainos atkurtas RMSE buvo {decimal(without_previous['Atkartotas_RMSE'], 4)}, "
        f"o R2 {decimal(without_previous['Atkartotas_R2'], 4)}. Modelis pagal kalendorinius "
        "požymius atkūrė bendrą ilgalaikį kainos kitimą, tačiau silpniau reagavo į staigius "
        "šuolius. Rezultatas nuo straipsnio skyrėsi dėl neaprašytų 81 įrašo pašalinimo, "
        "mokymo ir testavimo skirstymo bei nepateiktų optimizavimo detalių."
    )
    add_figure(
        doc,
        FIGURES / "tikrosios_ir_prognozuotos_kainos.png",
        "3 pav Tikrosios ir ANFIS prognozuotos testavimo kainos",
        5.80,
    )
    add_body(
        doc,
        f"Pridėjus ankstesnės savaitės kainą, RMSE sumažėjo iki "
        f"{decimal(with_previous['Atkartotas_RMSE'], 4)}, R2 padidėjo iki "
        f"{decimal(with_previous['Atkartotas_R2'], 4)}, o koreliacija pasiekė "
        f"{decimal(with_previous['Atkartota_koreliacija'], 4)}. Straipsnyje atitinkamai "
        "paskelbta 0,0532, 0,9970 ir 0,9985. Pagrindinė straipsnio išvada, kad ankstesnė "
        "kaina labai pagerina prognozę, buvo atkartota."
    )

    doc.add_page_break()
    add_heading(doc, "6 Ateities kainos prognozė")
    add_body(
        doc,
        "Papildomam bandymui modelis su ankstesnės savaitės kaina apmokytas su visa turima "
        "EIA istorija. Kiekvienos prognozuojamos savaitės rezultatas naudotas kaip kitos "
        "savaitės įvestis. 52 savaičių prognozė svyravo nuo "
        f"{decimal(min(year_values), 3)} iki {decimal(max(year_values), 3)} EUR už litrą."
    )
    add_figure(
        doc,
        FIGURES / "paskutiniai_3_metai_ir_metu_prognoze.png",
        "4 pav Paskutinių trejų metų reali kaina ir 52 savaičių prognozė",
        6.65,
    )
    add_body(
        doc,
        f"Ilgalaikę prognozę sudaro {len(forecast_2030)} savaitiniai taškai. Jos reikšmės "
        f"svyravo nuo {decimal(min(long_values), 3)} iki {decimal(max(long_values), 3)} EUR "
        f"už litrą, o paskutinė reikšmė buvo {decimal(long_values[-1], 3)} EUR už litrą. "
        "Ilgėjant rekursinei prognozei neapibrėžtumas kaupiasi, todėl šis grafikas rodo "
        "modelio ekstrapoliaciją, o ne garantuotą būsimą kuro kainą."
    )
    add_figure(
        doc,
        FIGURES / "tikroji_kaina_ir_prognoze_iki_2030.png",
        "5 pav Tikroji kaina ir ANFIS prognozė iki 2030 m.",
        6.20,
    )

    add_heading(doc, "7 Mokymo eiga")
    add_figure(
        doc,
        FIGURES / "mokymo_klaida.png",
        "6 pav ANFIS mokymo klaida logaritminėje skalėje",
        4.60,
    )
    add_body(
        doc,
        "Abiejų modelių mokymo klaida mažėjo ir stabilizavosi iki 100 epochos. Modelio su "
        "ankstesne kaina klaida viso mokymo metu buvo maždaug dviem dydžio eilėmis mažesnė. "
        "Tai patvirtina, kad artimiausia ankstesnė kaina turi daugiau prognozinės informacijos "
        "nei vien kalendoriniai požymiai."
    )

    add_heading(doc, "8 Atkūrimo ribotumai")
    add_body(
        doc,
        "Straipsnyje pateiktos MSE ir RMSE poros nėra matematiškai suderintos. Pavyzdžiui, "
        "RMSE 0,0532 kvadratas yra apie 0,00283, o paskelbtas MSE yra 0,04164. Analogiškas "
        "neatitikimas yra ir modelio be ankstesnės kainos lentelėje. Atkurtame eksperimente "
        "abi metrikos apskaičiuotos iš tų pačių prognozės liekanų, todėl tarpusavyje sutampa."
    )
    add_body(
        doc,
        "Atkūrimą taip pat riboja nepateiktas originalus programinis kodas, tikslūs atsitiktinės "
        "sekos parametrai ir nepaaiškintas 81 įrašo skirtumas. Dėl šių priežasčių vertinamas ne "
        "vien absoliutus metrikų sutapimas, bet ir pagrindinės išvados bei modelių santykis."
    )

    add_heading(doc, "9 Išvados")
    add_bullet(
        doc,
        "Savarankiška MATLAB ANFIS realizacija sėkmingai apmokyta be papildomų toolboxų.",
    )
    add_bullet(
        doc,
        "Modelis su ankstesnės savaitės kaina beveik tiksliai atkartojo straipsnio RMSE, R2 ir koreliaciją.",
    )
    add_bullet(
        doc,
        "Ankstesnė kaina sumažino atkurtą RMSE nuo 0,3682 iki 0,0521 USD už galoną.",
    )
    add_bullet(
        doc,
        "Papildoma prognozė parodė 52 savaičių ir iki 2030 m. modelio ekstrapoliaciją EUR už litrą.",
    )
    add_bullet(
        doc,
        "Straipsnio imties dydžio ir MSE bei RMSE neatitikimai neleidžia visiškai tiksliai atkurti visų rezultatų.",
    )

    add_heading(doc, "Literatūra")
    add_body(
        doc,
        "Eliwa E H I, El Koshiry A M, Abd El Hafeez T, Omar A. Optimal Gasoline Price "
        "Predictions Leveraging the ANFIS Regression Model. International Journal of "
        "Intelligent Systems, 8462056. DOI 10.1155/2024/8462056."
    )
    add_body(
        doc,
        "U S Energy Information Administration. Weekly U S All Grades All Formulations "
        "Retail Gasoline Prices. Series EMM EPM0 PTE NUS DPG. Jang J S R. ANFIS Adaptive "
        "Network Based Fuzzy Inference System. IEEE Transactions "
        "on Systems Man and Cybernetics, 23, 665-685."
    )

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build_report())
